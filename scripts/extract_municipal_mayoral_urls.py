#!/usr/bin/env python3
"""Extract mayoral URL claims from saved official municipal records.

The adapter reads local files only. It never makes a network request. One
manifest source can contain several files, but the adapter parses each file
once and reuses its records for every listed contest.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Mapping
from urllib.parse import urlsplit, urlunsplit
from xml.etree import ElementTree

import pandas as pd
from bs4 import BeautifulSoup


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.mayoral_discovery import normalize_person_name, unwrap_wayback_url


try:
    from pypdf import PdfReader as _PdfReader
except ImportError:  # The attempt ledger reports this optional dependency.
    _PdfReader = None

PdfReader = _PdfReader

PARSER_VERSION = "official_municipal_records_v1"
TERMINAL_STATUSES = {
    "found",
    "no_result",
    "source_unavailable",
    "request_error",
    "match_unresolved",
}

EVIDENCE_COLUMNS = [
    "claim_id",
    "source_id",
    "parser_version",
    "candidate_year_id",
    "candidate_contest_id",
    "contest_id",
    "ledb_candid",
    "full_name",
    "geo_name",
    "state_abb",
    "year",
    "month",
    "election_round",
    "candidate_source",
    "candidate_source_doi",
    "source_name",
    "source_profile_url",
    "source_date",
    "local_source_file",
    "source_file_type",
    "source_record_id",
    "source_record_locator",
    "raw_url",
    "exact_normalized_url",
    "evidence_url",
    "site_root_url",
    "registrable_domain",
    "url_type",
    "source_candidate_name",
    "identity_match",
    "identity_match_score",
    "identity_match_method",
    "temporal_match",
    "evidence_grade",
    "review_status",
    "acceptance",
    "evidence_text",
    "extraction_method",
    "retrieved_on",
]

ATTEMPT_COLUMNS = [
    "attempt_id",
    "source_id",
    "parser_version",
    "contest_id",
    "geo_name",
    "state_abb",
    "year",
    "month",
    "election_round",
    "source_name",
    "source_profile_url",
    "source_date",
    "local_files",
    "files_expected",
    "files_parsed",
    "url_records_found",
    "matched_claims",
    "unmatched_url_records",
    "source_status",
    "error",
    "attempted_on",
]

URL_PATTERN = re.compile(
    r"(?i)\b(?:https?://|www\.)[^\s<>\"'\]\[{}]+"
)
BARE_DOMAIN_PATTERN = re.compile(
    r"(?i)^(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+"
    r"(?:com|org|net|us|info|biz|co|me|io|gov)(?:[/:?#][^\s]*)?$"
)
URL_FIELD_PATTERN = re.compile(
    r"(?i)(?:^|[_\s-])(url|website|web|homepage|home_page|campaign_site|site)(?:$|[_\s-])"
)
TRAILING_URL_PUNCTUATION = ".,;:!?)]}"

SOCIAL_HOSTS = {
    "facebook.com",
    "instagram.com",
    "linkedin.com",
    "tiktok.com",
    "twitter.com",
    "x.com",
    "youtube.com",
}
FUNDRAISING_HOSTS = {
    "actblue.com",
    "anedot.com",
    "donorbox.org",
    "gofundme.com",
    "paypal.com",
    "winred.com",
}
SHARED_HOSTS = {
    "blogspot.com",
    "google.com",
    "nationbuilder.com",
    "sites.google.com",
    "squarespace.com",
    "tumblr.com",
    "weebly.com",
    "wixsite.com",
    "wordpress.com",
}
COMMON_SECOND_LEVEL_SUFFIXES = {
    "co.uk",
    "com.au",
    "com.br",
    "com.mx",
    "org.uk",
    "org.au",
}


class SourceUnavailableError(RuntimeError):
    """A listed local source cannot be read in the current environment."""


class SourceParseError(RuntimeError):
    """A listed local source exists but its content cannot be parsed."""


@dataclass(frozen=True)
class SourceRecord:
    """One local record that contains one URL and its nearby source text."""

    raw_url: str
    context: str
    locator: str
    extraction_method: str
    file_path: Path
    file_type: str


@dataclass(frozen=True)
class FileResult:
    """The parse result for one local source file."""

    file_path: Path
    file_type: str
    records: tuple[SourceRecord, ...]


def stable_id(prefix: str, *values: object) -> str:
    payload = "\x1f".join(str(value) for value in values)
    digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()[:24]
    return f"{prefix}_{digest}"


def clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).split())


def extract_urls(value: object, allow_bare: bool = False) -> list[str]:
    """Return source URL strings without changing their path or case."""
    text = clean_text(value)
    found = [match.group(0).rstrip(TRAILING_URL_PUNCTUATION) for match in URL_PATTERN.finditer(text)]
    if allow_bare:
        stripped = text.strip().rstrip(TRAILING_URL_PUNCTUATION)
        if BARE_DOMAIN_PATTERN.fullmatch(stripped):
            found.append(stripped)
    return list(dict.fromkeys(item for item in found if item))


def normalize_exact_url(value: str) -> str:
    """Normalize the scheme and host while preserving path and query case."""
    raw = unwrap_wayback_url(str(value).strip().replace("&amp;", "&"))
    lowered = raw.casefold()
    if lowered.startswith("www.") or BARE_DOMAIN_PATTERN.fullmatch(raw):
        raw = f"http://{raw}"
        lowered = raw.casefold()
    if not lowered.startswith(("http://", "https://")):
        return ""
    parsed = urlsplit(raw)
    if not parsed.hostname:
        return ""
    scheme = parsed.scheme.casefold()
    host = parsed.hostname.casefold()
    if ":" in host and not host.startswith("["):
        host = f"[{host}]"
    port = f":{parsed.port}" if parsed.port else ""
    userinfo = ""
    if parsed.username:
        userinfo = parsed.username
        if parsed.password:
            userinfo += f":{parsed.password}"
        userinfo += "@"
    return urlunsplit((scheme, f"{userinfo}{host}{port}", parsed.path, parsed.query, ""))


def site_root_url(value: str) -> str:
    exact = normalize_exact_url(value)
    if not exact:
        return ""
    parsed = urlsplit(exact)
    return urlunsplit((parsed.scheme, parsed.netloc, "/", "", ""))


def registrable_domain(value: str) -> str:
    exact = normalize_exact_url(value)
    host = (urlsplit(exact).hostname or "").casefold().removeprefix("www.")
    if not host or re.fullmatch(r"\d+(?:\.\d+){3}", host) or ":" in host:
        return host
    labels = host.split(".")
    if len(labels) <= 2:
        return host
    suffix = ".".join(labels[-2:])
    if suffix in COMMON_SECOND_LEVEL_SUFFIXES and len(labels) >= 3:
        return ".".join(labels[-3:])
    return suffix


def host_matches(host: str, domains: Iterable[str]) -> bool:
    return any(host == domain or host.endswith(f".{domain}") for domain in domains)


def classify_url(value: str) -> str:
    exact = normalize_exact_url(value)
    parsed = urlsplit(exact)
    host = (parsed.hostname or "").casefold().removeprefix("www.")
    path_has_identity = parsed.path not in {"", "/"}
    if host_matches(host, SOCIAL_HOSTS):
        return "social_profile"
    if host_matches(host, FUNDRAISING_HOSTS):
        return "fundraising_profile"
    if host.endswith(".gov"):
        return "officeholder_or_government_profile"
    if host_matches(host, SHARED_HOSTS):
        return "shared_host_path" if path_has_identity else "shared_host_root"
    registered = registrable_domain(exact)
    bare_host = host.removeprefix("www.")
    if registered and bare_host != registered:
        return "candidate_subdomain"
    return "custom_domain"


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceParseError(f"No supported text encoding: {path}")


def make_record(
    raw_url: str,
    context: object,
    locator: str,
    method: str,
    path: Path,
    file_type: str,
) -> SourceRecord | None:
    raw_url = str(raw_url).strip().rstrip(TRAILING_URL_PUNCTUATION)
    if not normalize_exact_url(raw_url):
        return None
    return SourceRecord(
        raw_url=raw_url,
        context=clean_text(context)[:2000],
        locator=locator,
        extraction_method=method,
        file_path=path,
        file_type=file_type,
    )


def parse_html(path: Path) -> list[SourceRecord]:
    soup = BeautifulSoup(read_text(path), "html.parser")
    records: list[SourceRecord] = []
    for number, anchor in enumerate(soup.find_all("a", href=True), start=1):
        container = anchor.find_parent(["tr", "li", "p", "article", "section", "div"])
        context = (container or anchor.parent or anchor).get_text(" ", strip=True)
        record = make_record(
            anchor.get("href", ""), context, f"anchor:{number}",
            "html_hyperlink", path, "html",
        )
        if record:
            records.append(record)

    containers = soup.find_all(["tr", "li", "p"])
    if not containers:
        containers = [soup.body or soup]
    for number, container in enumerate(containers, start=1):
        context = container.get_text(" ", strip=True)
        for raw_url in extract_urls(context):
            record = make_record(
                raw_url, context, f"text_record:{number}",
                "html_visible_text", path, "html",
            )
            if record:
                records.append(record)
    return records


def records_from_mapping_row(
    row: Mapping[str, object], locator: str, path: Path, file_type: str
) -> list[SourceRecord]:
    context = " | ".join(
        f"{key}={clean_text(value)}" for key, value in row.items() if clean_text(value)
    )
    records: list[SourceRecord] = []
    for key, value in row.items():
        for raw_url in extract_urls(value, allow_bare=bool(URL_FIELD_PATTERN.search(str(key)))):
            record = make_record(
                raw_url, context, locator, f"{file_type}_field:{key}", path, file_type
            )
            if record:
                records.append(record)
    return records


def parse_delimited(path: Path, delimiter: str, file_type: str) -> list[SourceRecord]:
    try:
        rows = csv.DictReader(read_text(path).splitlines(), delimiter=delimiter)
        if not rows.fieldnames:
            raise SourceParseError(f"The {file_type.upper()} file has no header: {path}")
        records: list[SourceRecord] = []
        for number, row in enumerate(rows, start=2):
            records.extend(records_from_mapping_row(row, f"row:{number}", path, file_type))
        return records
    except csv.Error as error:
        raise SourceParseError(f"Invalid {file_type.upper()} file {path}: {error}") from error


def iter_json_rows(value: object, locator: str = "root") -> Iterable[tuple[str, Mapping[str, object]]]:
    if isinstance(value, list):
        for index, item in enumerate(value):
            yield from iter_json_rows(item, f"{locator}[{index}]")
        return
    if not isinstance(value, dict):
        return
    scalars: dict[str, object] = {}
    for key, item in value.items():
        if isinstance(item, (str, int, float, bool)) or item is None:
            scalars[str(key)] = item
        elif isinstance(item, list) and all(
            isinstance(entry, (str, int, float, bool)) or entry is None for entry in item
        ):
            scalars[str(key)] = " | ".join(clean_text(entry) for entry in item)
    if scalars:
        yield locator, scalars
    for key, item in value.items():
        if isinstance(item, (dict, list)):
            yield from iter_json_rows(item, f"{locator}.{key}")


def parse_json(path: Path) -> list[SourceRecord]:
    try:
        value = json.loads(read_text(path))
    except json.JSONDecodeError as error:
        raise SourceParseError(f"Invalid JSON file {path}: {error}") from error
    records: list[SourceRecord] = []
    for locator, row in iter_json_rows(value):
        records.extend(records_from_mapping_row(row, locator, path, "json"))
    return records


def excel_column_index(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference.upper())
    if not letters:
        return 0
    result = 0
    for letter in letters.group(0):
        result = result * 26 + ord(letter) - ord("A") + 1
    return result - 1


def xlsx_cell_text(cell: ElementTree.Element, shared: list[str], namespace: str) -> str:
    cell_type = cell.get("t", "")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.findall(f".//{{{namespace}}}t"))
    value = cell.find(f"{{{namespace}}}v")
    if value is None or value.text is None:
        return ""
    if cell_type == "s":
        try:
            return shared[int(value.text)]
        except (IndexError, ValueError):
            return ""
    return value.text


def parse_xlsx_without_openpyxl(path: Path) -> list[SourceRecord]:
    """Read standard XLSX rows and hyperlinks with the Python standard library."""
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    try:
        with zipfile.ZipFile(path) as workbook:
            shared: list[str] = []
            if "xl/sharedStrings.xml" in workbook.namelist():
                root = ElementTree.fromstring(workbook.read("xl/sharedStrings.xml"))
                for item in root.findall(f"{{{spreadsheet_ns}}}si"):
                    shared.append("".join(node.text or "" for node in item.findall(f".//{{{spreadsheet_ns}}}t")))

            records: list[SourceRecord] = []
            sheet_names = sorted(
                name for name in workbook.namelist()
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
            )
            for sheet_name in sheet_names:
                sheet = ElementTree.fromstring(workbook.read(sheet_name))
                relationship_targets: dict[str, str] = {}
                rel_name = (
                    "xl/worksheets/_rels/" + Path(sheet_name).name + ".rels"
                )
                if rel_name in workbook.namelist():
                    rel_root = ElementTree.fromstring(workbook.read(rel_name))
                    for relation in rel_root.findall(f"{{{package_rel_ns}}}Relationship"):
                        relationship_targets[relation.get("Id", "")] = relation.get("Target", "")

                hyperlinks: dict[str, str] = {}
                for link in sheet.findall(f".//{{{spreadsheet_ns}}}hyperlink"):
                    relation_id = link.get(f"{{{office_rel_ns}}}id", "")
                    if relation_id in relationship_targets:
                        hyperlinks[link.get("ref", "")] = relationship_targets[relation_id]

                rows = sheet.findall(f".//{{{spreadsheet_ns}}}row")
                headers: list[str] = []
                for row_number, row in enumerate(rows, start=1):
                    cells: dict[int, tuple[str, str]] = {}
                    for cell in row.findall(f"{{{spreadsheet_ns}}}c"):
                        reference = cell.get("r", "")
                        cells[excel_column_index(reference)] = (
                            xlsx_cell_text(cell, shared, spreadsheet_ns), reference
                        )
                    width = max(cells, default=-1) + 1
                    values = [cells.get(index, ("", ""))[0] for index in range(width)]
                    if row_number == 1:
                        headers = [value or f"column_{index + 1}" for index, value in enumerate(values)]
                        continue
                    if len(headers) < width:
                        headers.extend(f"column_{index + 1}" for index in range(len(headers), width))
                    mapping = {headers[index]: values[index] for index in range(width)}
                    locator = f"{Path(sheet_name).stem}!row:{row_number}"
                    records.extend(records_from_mapping_row(mapping, locator, path, "xlsx"))
                    context = " | ".join(
                        f"{headers[index]}={values[index]}" for index in range(width) if values[index]
                    )
                    for _, (_, reference) in cells.items():
                        target = hyperlinks.get(reference, "")
                        if target:
                            record = make_record(
                                target, context, locator,
                                "xlsx_hyperlink", path, "xlsx",
                            )
                            if record:
                                records.append(record)
            return records
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as error:
        raise SourceParseError(f"Invalid XLSX file {path}: {error}") from error


def parse_xlsx(path: Path) -> list[SourceRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return parse_xlsx_without_openpyxl(path)

    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
        records: list[SourceRecord] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows())
            if not rows:
                continue
            headers = [clean_text(cell.value) or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
            for row_number, row in enumerate(rows[1:], start=2):
                mapping = {headers[index]: cell.value for index, cell in enumerate(row)}
                locator = f"{sheet.title}!row:{row_number}"
                records.extend(records_from_mapping_row(mapping, locator, path, "xlsx"))
                context = " | ".join(
                    f"{headers[index]}={clean_text(cell.value)}"
                    for index, cell in enumerate(row) if clean_text(cell.value)
                )
                for cell in row:
                    if cell.hyperlink and cell.hyperlink.target:
                        record = make_record(
                            cell.hyperlink.target, context, locator,
                            "xlsx_hyperlink", path, "xlsx",
                        )
                        if record:
                            records.append(record)
        return records
    except Exception as error:
        raise SourceParseError(f"Invalid XLSX file {path}: {error}") from error


def pdf_object(value: object) -> object:
    return value.get_object() if hasattr(value, "get_object") else value


def parse_pdf(path: Path) -> list[SourceRecord]:
    if PdfReader is None:
        raise SourceUnavailableError(
            "PDF extraction requires the optional pypdf package"
        )
    try:
        reader = PdfReader(str(path))
        records: list[SourceRecord] = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = clean_text(page.extract_text() or "")
            for raw_url in extract_urls(page_text):
                record = make_record(
                    raw_url, page_text, f"page:{page_number}",
                    "pdf_visible_text", path, "pdf",
                )
                if record:
                    records.append(record)
            for annotation_reference in page.get("/Annots") or []:
                annotation = pdf_object(annotation_reference)
                if not isinstance(annotation, Mapping) or str(annotation.get("/Subtype", "")) != "/Link":
                    continue
                action = pdf_object(annotation.get("/A", {}))
                raw_url = action.get("/URI", "") if isinstance(action, Mapping) else ""
                if not raw_url:
                    continue
                context = clean_text(annotation.get("/Contents", "")) or page_text
                record = make_record(
                    str(raw_url), context, f"page:{page_number}:annotation",
                    "pdf_hyperlink_annotation", path, "pdf",
                )
                if record:
                    records.append(record)
        return records
    except SourceUnavailableError:
        raise
    except Exception as error:
        raise SourceParseError(f"Invalid PDF file {path}: {error}") from error


def parse_local_file(path: Path, file_type: str | None = None) -> FileResult:
    if str(path).startswith(("http://", "https://")):
        raise SourceUnavailableError("The adapter accepts local files only")
    if not path.exists() or not path.is_file():
        raise SourceUnavailableError(f"Local source file is unavailable: {path}")
    kind = (file_type or path.suffix.removeprefix(".")).casefold()
    parsers = {
        "html": parse_html,
        "htm": parse_html,
        "csv": lambda item: parse_delimited(item, ",", "csv"),
        "tsv": lambda item: parse_delimited(item, "\t", "tsv"),
        "xlsx": parse_xlsx,
        "json": parse_json,
        "pdf": parse_pdf,
    }
    if kind not in parsers:
        raise SourceUnavailableError(f"Unsupported municipal source type: {kind or '(none)'}")
    records = parsers[kind](path)
    return FileResult(path, "html" if kind == "htm" else kind, tuple(records))


def resolve_local_path(value: str, manifest_path: Path) -> Path:
    if str(value).startswith(("http://", "https://")):
        return Path(str(value))
    path = Path(value).expanduser()
    if path.is_absolute():
        return path
    manifest_relative = manifest_path.parent / path
    if manifest_relative.exists():
        return manifest_relative.resolve()
    return (Path.cwd() / path).resolve()


def read_universe(path: Path) -> pd.DataFrame:
    data = path.read_bytes()
    frame = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            frame = pd.read_csv(path, encoding=encoding, low_memory=False)
            break
        except UnicodeDecodeError:
            continue
    if frame is None:
        raise ValueError(f"The candidate universe has no supported encoding: {path}")
    required = {
        "candidate_year_id", "candidate_contest_id", "ledb_candid", "full_name",
        "geo_name", "state_abb", "year", "month",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Candidate universe missing columns: {sorted(missing)}")
    if "contest_id" not in frame.columns and "contest" not in frame.columns:
        raise ValueError("Candidate universe requires contest_id or contest")
    if frame.duplicated("candidate_contest_id").any():
        raise ValueError("Candidate universe has duplicate candidate_contest_id values")
    return frame


def normalized_equals(series: pd.Series, value: object) -> pd.Series:
    expected = clean_text(value).casefold()
    return series.fillna("").astype(str).str.strip().str.casefold().eq(expected)


def source_target_rows(universe: pd.DataFrame, source: Mapping[str, Any]) -> pd.DataFrame:
    target = universe[
        normalized_equals(universe["geo_name"], source["geo_name"])
        & normalized_equals(universe["state_abb"], source["state_abb"])
        & pd.to_numeric(universe["year"], errors="coerce").eq(int(source["year"]))
    ].copy()
    if "in_discovery_scope" in target.columns and not source.get("include_out_of_scope", False):
        scope = target["in_discovery_scope"]
        if scope.dtype == bool:
            target = target[scope]
        else:
            target = target[scope.astype(str).str.casefold().isin({"true", "1", "yes"})]
    if source.get("month") not in (None, ""):
        target = target[pd.to_numeric(target["month"], errors="coerce").eq(int(source["month"]))]
    contest_column = "contest_id" if "contest_id" in target.columns else "contest"
    contest_values = source.get("contest_ids")
    if contest_values is None and source.get("contest_id") not in (None, ""):
        contest_values = [source["contest_id"]]
    if contest_values is not None:
        wanted = {str(value) for value in contest_values}
        target = target[target[contest_column].astype(str).isin(wanted)]
    return target


def candidate_match(context: str, candidates: pd.DataFrame) -> tuple[pd.Series | None, float, str]:
    normalized_context = f" {normalize_person_name(context)} "
    exact_matches: list[int] = []
    token_matches: list[int] = []
    for index, candidate in candidates.reset_index(drop=True).iterrows():
        normalized_name = normalize_person_name(candidate["full_name"])
        if normalized_name and f" {normalized_name} " in normalized_context:
            exact_matches.append(index)
            continue
        tokens = normalized_name.split()
        if len(tokens) >= 2:
            first_last = f" {tokens[0]} {tokens[-1]} "
            reversed_name = f" {tokens[-1]} {tokens[0]} "
            if first_last in normalized_context or reversed_name in normalized_context:
                token_matches.append(index)
    reset = candidates.reset_index(drop=True)
    if len(exact_matches) == 1:
        return reset.iloc[exact_matches[0]], 1.0, "unique_exact_full_name"
    if not exact_matches and len(token_matches) == 1:
        return reset.iloc[token_matches[0]], 0.90, "unique_first_last_name"
    return None, 0.0, "ambiguous_or_missing_name"


def election_round(source: Mapping[str, Any], candidates: pd.DataFrame) -> str:
    supplied = clean_text(source.get("election_round", ""))
    if supplied:
        return supplied
    for column in ("election_round", "election_stage_status", "stage_scope_status"):
        if column in candidates.columns:
            values = [clean_text(value) for value in candidates[column].dropna().unique() if clean_text(value)]
            if len(values) == 1:
                return values[0]
    return "unresolved"


def deduplicate_records(records: Iterable[SourceRecord]) -> list[SourceRecord]:
    """Remove parser duplicates within a file while preserving each source file claim."""
    by_key: dict[tuple[str, str, str], SourceRecord] = {}
    methods: dict[tuple[str, str, str], list[str]] = {}
    for record in records:
        exact = normalize_exact_url(record.raw_url)
        context_key = normalize_person_name(record.context)
        key = (str(record.file_path), exact, context_key)
        methods.setdefault(key, []).append(record.extraction_method)
        current = by_key.get(key)
        if current is None or len(record.context) > len(current.context):
            by_key[key] = record
    result = []
    for key, record in by_key.items():
        result.append(SourceRecord(
            raw_url=record.raw_url,
            context=record.context,
            locator=record.locator,
            extraction_method="|".join(dict.fromkeys(methods[key])),
            file_path=record.file_path,
            file_type=record.file_type,
        ))
    return sorted(result, key=lambda item: (str(item.file_path), normalize_exact_url(item.raw_url)))


def validate_source(source: Mapping[str, Any]) -> None:
    required = {"source_id", "geo_name", "state_abb", "year", "local_files"}
    missing = required - set(source)
    if missing:
        raise ValueError(f"Municipal source missing fields: {sorted(missing)}")
    if not re.fullmatch(r"[a-z0-9][a-z0-9_.-]*", str(source["source_id"])):
        raise ValueError(f"Invalid stable source_id: {source['source_id']}")
    if not isinstance(source["local_files"], list):
        raise ValueError(f"Source {source['source_id']} requires a local_files list")
    upstream_status = clean_text(source.get("upstream_status", ""))
    if not source["local_files"] and upstream_status not in {
        "source_unavailable", "request_error",
    }:
        raise ValueError(
            f"Source {source['source_id']} requires local files or a terminal upstream_status"
        )


def build_attempt(
    source: Mapping[str, Any], contest_id: str, candidates: pd.DataFrame,
    file_specs: list[tuple[Path, str | None, str]], parsed: list[FileResult],
    unavailable_errors: list[str], parse_errors: list[str], records: list[SourceRecord],
    evidence: list[dict[str, Any]], unmatched: int,
) -> dict[str, Any]:
    round_name = election_round(source, candidates)
    months = pd.to_numeric(candidates.get("month", pd.Series(dtype=float)), errors="coerce").dropna().unique()
    month = int(months[0]) if len(months) == 1 else source.get("month", "")
    if evidence:
        status = "found"
    elif unavailable_errors:
        status = "source_unavailable"
    elif parse_errors:
        status = "request_error"
    elif records:
        status = "match_unresolved"
    else:
        status = "no_result"
    if status not in TERMINAL_STATUSES:
        raise AssertionError(f"Invalid terminal status: {status}")
    errors = unavailable_errors + parse_errors
    return {
        "attempt_id": stable_id("source_attempt", source["source_id"], contest_id),
        "source_id": source["source_id"],
        "parser_version": PARSER_VERSION,
        "contest_id": contest_id,
        "geo_name": source["geo_name"],
        "state_abb": source["state_abb"],
        "year": int(source["year"]),
        "month": month,
        "election_round": round_name,
        "source_name": source.get("source_name", source["source_id"]),
        "source_profile_url": source.get("source_url", ""),
        "source_date": source.get("source_date", ""),
        "local_files": "|".join(file_key for _, _, file_key in file_specs),
        "files_expected": len(file_specs),
        "files_parsed": len(parsed),
        "url_records_found": len(records),
        "matched_claims": len(evidence),
        "unmatched_url_records": unmatched,
        "source_status": status,
        "error": " | ".join(errors)[:2000],
        "attempted_on": date.today().isoformat(),
    }


def extract_manifest(manifest_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Extract evidence and one terminal attempt for each contest-source target."""
    manifest_path = manifest_path.resolve()
    manifest = json.loads(read_text(manifest_path))
    if not isinstance(manifest, dict) or "candidate_universe" not in manifest:
        raise ValueError("Manifest requires candidate_universe")
    if not isinstance(manifest.get("sources"), list):
        raise ValueError("Manifest requires a sources list")
    universe_path = resolve_local_path(str(manifest["candidate_universe"]), manifest_path)
    universe = read_universe(universe_path)
    evidence_rows: list[dict[str, Any]] = []
    attempt_rows: list[dict[str, Any]] = []
    parse_cache: dict[tuple[Path, str | None], tuple[str, object]] = {}

    seen_source_ids: set[str] = set()
    for source in manifest["sources"]:
        if not isinstance(source, dict):
            raise ValueError("Each municipal source must be an object")
        validate_source(source)
        if source["source_id"] in seen_source_ids:
            raise ValueError(f"Duplicate source_id in manifest: {source['source_id']}")
        seen_source_ids.add(source["source_id"])

        file_specs: list[tuple[Path, str | None, str]] = []
        for item in source["local_files"]:
            if isinstance(item, str):
                file_specs.append((resolve_local_path(item, manifest_path), None, item))
            elif isinstance(item, dict) and "path" in item:
                raw_path = str(item["path"])
                file_specs.append((
                    resolve_local_path(raw_path, manifest_path),
                    clean_text(item.get("type")) or None,
                    clean_text(item.get("file_id")) or raw_path,
                ))
            else:
                raise ValueError(f"Invalid local file entry for {source['source_id']}")

        parsed: list[FileResult] = []
        unavailable_errors: list[str] = []
        parse_errors: list[str] = []
        upstream_status = clean_text(source.get("upstream_status", ""))
        upstream_error = clean_text(source.get("upstream_error", ""))
        if upstream_status == "source_unavailable":
            unavailable_errors.append(upstream_error or "Upstream source was unavailable")
        elif upstream_status == "request_error":
            parse_errors.append(upstream_error or "Upstream source request failed")
        for path, kind, _ in file_specs:
            cache_key = (path, kind)
            if cache_key not in parse_cache:
                try:
                    parse_cache[cache_key] = ("parsed", parse_local_file(path, kind))
                except SourceUnavailableError as error:
                    parse_cache[cache_key] = ("unavailable", str(error))
                except (SourceParseError, OSError, UnicodeError) as error:
                    parse_cache[cache_key] = ("error", str(error))
            cache_status, cached = parse_cache[cache_key]
            if cache_status == "parsed":
                parsed.append(cached)  # type: ignore[arg-type]
            elif cache_status == "unavailable":
                unavailable_errors.append(str(cached))
            else:
                parse_errors.append(str(cached))

        records = deduplicate_records(record for result in parsed for record in result.records)
        file_keys = {path: file_key for path, _, file_key in file_specs}
        targets = source_target_rows(universe, source)
        contest_column = "contest_id" if "contest_id" in targets.columns else "contest"
        target_groups: list[tuple[str, pd.DataFrame]] = [
            (str(contest_id), group.copy())
            for contest_id, group in targets.groupby(contest_column, dropna=False, sort=True)
        ]
        if not target_groups:
            fallback_id = str(source.get("contest_id") or stable_id(
                "unresolved_contest", source["geo_name"], source["state_abb"], source["year"], source.get("month", "")
            ))
            empty_candidates = universe.iloc[0:0].copy()
            attempt = build_attempt(
                source, fallback_id, empty_candidates, file_specs, parsed,
                unavailable_errors, parse_errors, records, [], len(records),
            )
            attempt["source_status"] = (
                "source_unavailable" if unavailable_errors else
                "request_error" if parse_errors else "match_unresolved"
            )
            missing_message = "No candidate universe rows matched this source target"
            attempt["error"] = " | ".join(item for item in [attempt["error"], missing_message] if item)
            attempt_rows.append(attempt)
            continue

        for contest_id, candidates in target_groups:
            contest_evidence: list[dict[str, Any]] = []
            unmatched = 0
            for record in records:
                candidate, score, method = candidate_match(record.context, candidates)
                if candidate is None:
                    unmatched += 1
                    continue
                exact = normalize_exact_url(record.raw_url)
                record_id = stable_id(
                    "source_record", source["source_id"],
                    file_keys.get(record.file_path, record.file_path.name), record.locator
                )
                claim_id = stable_id(
                    "municipal_claim", source["source_id"],
                    file_keys.get(record.file_path, record.file_path.name),
                    contest_id, candidate["candidate_contest_id"], exact,
                )
                source_candidate_name = clean_text(candidate["full_name"])
                row = {
                    "claim_id": claim_id,
                    "source_id": source["source_id"],
                    "parser_version": PARSER_VERSION,
                    "candidate_year_id": candidate["candidate_year_id"],
                    "candidate_contest_id": candidate["candidate_contest_id"],
                    "contest_id": contest_id,
                    "ledb_candid": candidate["ledb_candid"],
                    "full_name": candidate["full_name"],
                    "geo_name": candidate["geo_name"],
                    "state_abb": candidate["state_abb"],
                    "year": int(candidate["year"]),
                    "month": int(candidate["month"]) if not pd.isna(candidate["month"]) else "",
                    "election_round": election_round(source, candidates),
                    "candidate_source": candidate.get("source_dataset", "American Local Government Elections Database"),
                    "candidate_source_doi": candidate.get("source_doi", "10.17605/OSF.IO/MV5E6"),
                    "source_name": source.get("source_name", source["source_id"]),
                    "source_profile_url": source.get("source_url", ""),
                    "source_date": source.get("source_date", ""),
                    "local_source_file": file_keys.get(record.file_path, record.file_path.name),
                    "source_file_type": record.file_type,
                    "source_record_id": record_id,
                    "source_record_locator": record.locator,
                    "raw_url": record.raw_url,
                    "exact_normalized_url": exact,
                    "evidence_url": exact,
                    "site_root_url": site_root_url(exact),
                    "registrable_domain": registrable_domain(exact),
                    "url_type": classify_url(exact),
                    "source_candidate_name": source_candidate_name,
                    "identity_match": "matched",
                    "identity_match_score": score,
                    "identity_match_method": method,
                    "temporal_match": "pending_archive_validation",
                    "evidence_grade": "C",
                    "review_status": "pending",
                    "acceptance": "",
                    "evidence_text": record.context[:500],
                    "extraction_method": record.extraction_method,
                    "retrieved_on": date.today().isoformat(),
                }
                contest_evidence.append(row)

            # A file can repeat one claim. Keep separate claims across source files.
            by_claim = {row["claim_id"]: row for row in contest_evidence}
            contest_evidence = sorted(by_claim.values(), key=lambda row: row["claim_id"])
            evidence_rows.extend(contest_evidence)
            attempt_rows.append(build_attempt(
                source, contest_id, candidates, file_specs, parsed,
                unavailable_errors, parse_errors, records, contest_evidence, unmatched,
            ))

    evidence = pd.DataFrame(evidence_rows).reindex(columns=EVIDENCE_COLUMNS)
    attempts = pd.DataFrame(attempt_rows).reindex(columns=ATTEMPT_COLUMNS)
    if not attempts.empty:
        if attempts.duplicated(["source_id", "contest_id"]).any():
            raise ValueError("Attempt ledger has duplicate contest-source targets")
        invalid = set(attempts["source_status"]) - TERMINAL_STATUSES
        if invalid:
            raise ValueError(f"Attempt ledger has invalid terminal statuses: {sorted(invalid)}")
    if not evidence.empty and evidence.duplicated("claim_id").any():
        raise ValueError("Evidence queue has duplicate claim_id values")
    return evidence, attempts


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    frame.to_csv(temporary, index=False)
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--evidence-output", required=True, type=Path)
    parser.add_argument("--attempt-output", required=True, type=Path)
    args = parser.parse_args()
    evidence, attempts = extract_manifest(args.manifest)
    write_csv(evidence, args.evidence_output)
    write_csv(attempts, args.attempt_output)
    print({
        "evidence_rows": len(evidence),
        "attempt_rows": len(attempts),
        "terminal_statuses": attempts["source_status"].value_counts().sort_index().to_dict(),
    })


if __name__ == "__main__":
    main()
