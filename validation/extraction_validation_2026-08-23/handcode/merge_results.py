#!/usr/bin/env python3
"""Merge handcode results back into the validation workbook.

Usage:
  python3 merge_results.py cwe_record_precision_results.json cwe_page_recall_results.json \
      [-o candidate_extraction_validation_coded.xlsx]

Reads ../candidate_extraction_validation.xlsx, fills the coding columns from the two
exported results files (matching on validation_id / parent_document_id), and writes the
coded workbook. Refuses to run if any item is missing or not complete.
"""
import argparse
import json
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent / "candidate_extraction_validation.xlsx"

PRECISION_QS = ["record_present_in_page", "normalized_value_faithful",
                "record_type_correct", "evidence_alone_supports_record"]


def load_results(path, expected_task_id):
    res = json.loads(Path(path).read_text())
    if res.get("task_id") != expected_task_id:
        raise SystemExit(f"{path}: task_id {res.get('task_id')!r}, expected {expected_task_id!r}")
    by_id = {r["item_id"]: r for r in res["responses"]}
    incomplete = [i for i, r in by_id.items() if r.get("status") != "complete"]
    if incomplete:
        raise SystemExit(f"{path}: {len(incomplete)} items not complete (e.g. {incomplete[:3]})")
    return by_id


def fill(ws, key_col, mapping):
    header = {c.value: i + 1 for i, c in enumerate(ws[1])}
    n = 0
    for row in ws.iter_rows(min_row=2):
        key = str(row[header[key_col] - 1].value)
        if key not in mapping:
            raise SystemExit(f"sheet {ws.title!r}: no coded result for {key_col}={key}")
        r = mapping[key]
        for col, value in r.items():
            row[header[col] - 1].value = value
        n += 1
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("precision_results")
    ap.add_argument("recall_results")
    ap.add_argument("-o", "--out", default=str(HERE.parent / "candidate_extraction_validation_coded.xlsx"))
    args = ap.parse_args()

    prec = load_results(args.precision_results, "cwe_record_precision")
    rec = load_results(args.recall_results, "cwe_page_recall")

    wb = openpyxl.load_workbook(WORKBOOK)

    prec_map = {}
    for item_id, r in prec.items():
        a = r["answers"]
        row = {q: a[q] for q in PRECISION_QS}
        row["coding_status"] = "complete"
        row["coder_notes"] = a.get("notes") or None
        prec_map[item_id] = row
    n1 = fill(wb["Record precision"], "validation_id", prec_map)

    rec_map = {}
    for item_id, r in rec.items():
        a = r["answers"]
        captured = a["all_requested_records_captured"]
        missed = a.get("missed_records", [])
        if captured == "no" and not missed:
            raise SystemExit(f"recall item {item_id}: answered 'no' but no missed records listed")
        rec_map[item_id] = {
            "all_requested_records_captured": captured,
            "missed_records_json": json.dumps(missed, ensure_ascii=False) if missed else None,
            "coding_status": "complete",
            "coder_notes": a.get("notes") or None,
        }
    n2 = fill(wb["Page recall"], "parent_document_id", rec_map)

    wb.save(args.out)
    print(f"merged {n1} precision + {n2} recall items -> {args.out}")


if __name__ == "__main__":
    main()
