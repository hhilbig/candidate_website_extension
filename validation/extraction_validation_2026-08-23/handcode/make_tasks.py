#!/usr/bin/env python3
"""Convert candidate_extraction_validation.xlsx into two handcode task specs.

Produces cwe_record_precision.json and cwe_page_recall.json in this directory.
Build the HTML apps with:
  python3 ~/.agents/skills/handcode/scripts/build_task.py cwe_record_precision.json
  python3 ~/.agents/skills/handcode/scripts/build_task.py cwe_page_recall.json
"""
import json
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent / "candidate_extraction_validation.xlsx"

STRICT_RULE = ("Code only what the supplied text establishes. Do not use outside knowledge.")


def sheet_rows(ws):
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(header, row))


def make_precision(wb):
    items = []
    for r in sheet_rows(wb["Record precision"]):
        evidence = str(r.get("evidence") or "")
        items.append({
            "id": str(r["validation_id"]),
            "label": f"{r['candidate_name']} ({r['year']}) · {r['subtype']}",
            "display": [
                {"label": "Candidate", "format": "badges", "content": {
                    "Name": r["candidate_name"], "Year": r["year"], "Party": r["party"],
                    "Route": r["route"], "Subtype": r["subtype"],
                    "Evidence exact match": str(r["evidence_exact_match"]),
                }},
                {"label": "Model record", "format": "json", "content": str(r["model_record_json"])},
                {"label": "Source context", "format": "text",
                 "content": str(r["source_context"] or ""),
                 "highlight": [evidence] if evidence else []},
            ],
        })
    task = {
        "task_id": "cwe_record_precision",
        "title": "Record precision — candidate extraction validation",
        "instructions": (
            "### Record precision\n"
            "75 sampled model records. For each, compare the model record with the source context "
            "(the claimed evidence is highlighted) and answer the four questions.\n"
            f"- {STRICT_RULE}\n"
            "- Be strict on the last question: for example, a candidate's name followed by "
            "\"for Senate\" does not by itself prove prior elected office.\n"
            "- Add a brief note for any `no`."
        ),
        "items": items,
        "questions": [
            {"id": "record_present_in_page", "label": "Record present in page?", "type": "yesno",
             "help": "Does the page actually state the claimed fact, policy position, or endorsement?"},
            {"id": "normalized_value_faithful", "label": "Normalized value faithful?", "type": "yesno",
             "help": "Does the model's normalized value preserve the source's meaning?"},
            {"id": "record_type_correct", "label": "Record type correct?", "type": "yesno",
             "help": "Is the record assigned to the right category (see Route / Subtype badges)?"},
            {"id": "evidence_alone_supports_record", "label": "Evidence alone supports record?", "type": "yesno",
             "help": "Does the quoted (highlighted) evidence, by itself, establish the complete record? Be strict."},
        ],
    }
    return task


def make_recall(wb):
    items = []
    for r in sheet_rows(wb["Page recall"]):
        items.append({
            "id": str(r["parent_document_id"]),
            "label": f"{r['candidate_name']} ({r['year']}) · {r['route']}",
            "display": [
                {"label": "Page", "format": "badges", "content": {
                    "Name": r["candidate_name"], "Year": r["year"], "Party": r["party"],
                    "Route": r["route"], "Segments": r["segment_count"],
                    "Parse failures": r["parse_failure_count"],
                }},
                {"label": "All model records for this page", "format": "json",
                 "content": str(r["model_records_json"])},
                {"label": "Full source text", "format": "text", "content": str(r["source_text"] or "")},
            ],
        })
    task = {
        "task_id": "cwe_page_recall",
        "title": "Page recall — candidate extraction validation",
        "instructions": (
            "### Page recall\n"
            "15 complete pages. Read the full source text, review everything in the model records, "
            "and decide whether the model captured all requested information.\n"
            "Requested information means:\n"
            "- Issues: every distinct, concrete policy statement.\n"
            "- Biography: occupation, education, elected office, public service, military service, "
            "local roots, family, and other concrete biographical facts.\n"
            "- Endorsements: every actual endorser, excluding generic supporter claims and calls "
            "for visitors to endorse.\n"
            f"- {STRICT_RULE}\n"
            "If at least one qualifying record was missed, answer `no` and add one Missed records "
            "row per omission."
        ),
        "items": items,
        "questions": [
            {"id": "all_requested_records_captured", "label": "All requested records captured?", "type": "yesno",
             "help": "yes = nothing qualifying was missed; no = at least one qualifying record was missed."},
            {"id": "missed_records", "label": "Missed records", "type": "structured", "required": True,
             "show_if": {"question": "all_requested_records_captured", "equals": "no"},
             "fields": [
                 {"id": "type", "label": "Type", "placeholder": "e.g. education"},
                 {"id": "value", "label": "Value", "placeholder": "e.g. Graduated from UCLA"},
                 {"id": "evidence", "label": "Evidence quote", "placeholder": "verbatim from the page"},
             ],
             "help": "One row per record the model should have extracted but did not."},
        ],
    }
    return task


def main():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    for task in (make_precision(wb), make_recall(wb)):
        out = HERE / f"{task['task_id']}.json"
        out.write_text(json.dumps(task, ensure_ascii=False, indent=1))
        print(f"wrote {out} ({len(task['items'])} items)")


if __name__ == "__main__":
    main()
